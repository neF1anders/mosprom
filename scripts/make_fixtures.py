import struct
from pathlib import Path

import numpy as np
import trimesh
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

out = Path("tests/fixtures")
out.mkdir(parents=True, exist_ok=True)
box = trimesh.creation.box(extents=[10, 20, 30])
box.export(out / "box.stl")
trimesh.Trimesh(box.vertices, box.faces[:-1], process=False).export(out / "open.stl")
trimesh.Trimesh(box.vertices, np.vstack([box.faces, box.faces[0]]), process=False).export(
    out / "duplicate.stl"
)
second = box.copy()
second.apply_translation([40, 0, 0])
trimesh.util.concatenate([box, second]).export(out / "shells.stl")
nan = bytearray((out / "box.stl").read_bytes())
struct.pack_into("<f", nan, 96, float("nan"))
(out / "nan.stl").write_bytes(nan)
(out / "empty.txt").write_bytes(b"")
(out / "spoof.pdf").write_text("not a PDF")
(out / "corrupt.pdf").write_bytes(b"%PDF-1.4\ninvalid")
(out / "task.txt").write_text(
    "Номинальный момент 25 Нм\nПиковый момент 75 Нм\nТемпература от 10 до 70 C", encoding="utf-8"
)
(out / "conflict.txt").write_text("TEST_CONFLICT\nДлина 80 мм\nДлина 100 мм", encoding="utf-8")
doc = Document()
doc.add_paragraph("Номинальный момент 25 Нм")
doc.add_paragraph("Пиковый момент 75 Нм")
doc.save(out / "task.docx")
writer = PdfWriter()
page = writer.add_blank_page(600, 800)
font = DictionaryObject(
    {
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    }
)
page[NameObject("/Resources")] = DictionaryObject(
    {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
)
stream = DecodedStreamObject()
stream.set_data(
    b"BT /F1 18 Tf 40 700 Td (Nominal torque 25 Nm, peak torque 75 Nm. Temperature 10 to 70 C.) Tj ET"
)
page[NameObject("/Contents")] = writer._add_object(stream)
writer.write(out / "text.pdf")
writer.encrypt("secret")
writer.write(out / "encrypted.pdf")
writer = PdfWriter()
for _ in range(51):
    writer.add_blank_page(100, 100)
writer.write(out / "too-many-pages.pdf")
image = Image.new("RGB", (1600, 600), "white")
draw = ImageDraw.Draw(image)
fonts = [
    Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
    Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
]
font = ImageFont.truetype(str(next(p for p in fonts if p.exists())), 48)
draw.text((50, 100), "Nominal torque 25 Nm\nPeak torque 75 Nm", font=font, fill="black")
image.save(out / "scan.pdf", "PDF")
image.save(out / "photo.png")
writer = PdfWriter()
writer.append(out / "text.pdf")
writer.append(out / "scan.pdf")
writer.write(out / "mixed.pdf")
print("Generated synthetic fixtures")
