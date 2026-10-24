import pytest

from src.backend.domain.errors import AppError
from src.backend.infrastructure.documents import inspect_document

pytestmark = pytest.mark.unit


@pytest.mark.parametrize("kind", ["txt", "md", "pdf", "docx", "png"])
def test_T01_empty_files_rejected(tmp_path, kind):
    path = tmp_path / ("empty." + kind)
    path.write_bytes(b"")
    with pytest.raises(AppError):
        inspect_document(path, kind)


@pytest.mark.parametrize("kind", ["pdf", "docx", "png", "jpg"])
def test_T01_spoofed_extensions_rejected(tmp_path, kind):
    path = tmp_path / ("spoof." + kind)
    path.write_text("not this file format")
    with pytest.raises(AppError):
        inspect_document(path, kind)


def test_text_preserves_source(tmp_path):
    path = tmp_path / "task.txt"
    path.write_text("Номинальный момент 25 Нм, пиковый 75 Нм", encoding="utf-8")
    result = inspect_document(path, "txt")
    assert "25" in result["fragments"][0]["text"]
    assert result["fragments"][0]["page"] == 1


@pytest.mark.parametrize("limit", [1, 2, 3, 4, 5])
def test_configured_page_limit_is_enforced(limit):
    from pathlib import Path

    from src.backend.domain.limits import InputLimits

    with pytest.raises(AppError) as exc:
        inspect_document(
            Path("tests/fixtures/too-many-pages.pdf"), "pdf", InputLimits(pdf_pages=limit)
        )
    assert exc.value.code == "page_limit"


def test_stl_is_not_a_document(tmp_path):
    path = tmp_path / "model.stl"
    path.write_bytes(b"solid x")
    with pytest.raises(AppError) as exc:
        inspect_document(path, "stl")
    assert exc.value.code == "unsupported_type"


def test_docx_drawings_become_pages(tmp_path):
    from docx import Document
    from PIL import Image

    picture = tmp_path / "drawing.png"
    Image.new("RGB", (300, 200), "white").save(picture)
    document = Document()
    document.add_paragraph("Номинальный момент 25 Нм")
    document.add_picture(str(picture))
    path = tmp_path / "task"
    document.save(path)
    result = inspect_document(path, "docx")
    assert [f["page"] for f in result["fragments"]] == [1, 2]
    assert result["fragments"][1]["image"] == result["artifacts"][0]
    assert (tmp_path / result["artifacts"][0]).is_file()
